"""Build docs/cost_model.xlsx: monthly running cost for 25-500 new enquiries on WhatsApp and Instagram.

    python scripts/build_cost_model.py            # AI-cost cells filled from out/benchmark/results.csv if present
    python scripts/build_cost_model.py --results path/to/results.csv

Every number is an editable input on the Inputs sheet; the other sheets are formulas.
"""

from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
GREEN = Font(name=FONT, color="008000")
BLACK = Font(name=FONT)
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
HEAD = Font(name=FONT, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
SECTION_FILL = PatternFill("solid", fgColor="D9E1F2")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Border(bottom=Side(style="thin", color="BFBFBF"))

USD = '$#,##0.00;($#,##0.00);"-"'
USD4 = '$#,##0.0000;($#,##0.0000);"-"'
USD0 = '$#,##0;($#,##0);"-"'
INT = '#,##0;(#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
HRS = '#,##0.0;(#,##0.0);"-"'

VOLUMES = [25, 50, 100, 150, 200, 250, 300, 400, 500]
SCENARIOS = [0.50, 1.00, 2.00, 3.00, 5.00]

WA_SOURCE = (
    "Meta WhatsApp Business Platform rate card, Egypt, USD, per delivered message (marketing $0.0644 from 1 Jan 2026; "
    "utility/authentication $0.0036, July 2026 card). From 1 Oct 2026 service messages are billed at the utility rate "
    "after 1,000 free per business phone number per month, and utility templates inside the 24h window are billable. "
    "developers.facebook.com/docs/whatsapp/pricing. Checked 3 Oct 2026 via secondary sources; confirm in WhatsApp Manager."
)
IG_SOURCE = (
    "Instagram Messaging API (Graph API): no per-message fee. Replies only within 24h of the user's last message; the "
    "HUMAN_AGENT tag (7 days) needs its own App Review and may not be used for automated messages."
)


def benchmark_averages(path: Path) -> dict[str, float] | None:
    """Average cost of each step type over successful benchmark runs."""
    if not path.exists():
        return None
    with path.open(newline="", encoding="utf-8") as fh:
        rows = [r for r in csv.DictReader(fh) if not r.get("error")]
    out = {}
    for key, prefix in (("intake", "intake"), ("draft", "draft"), ("revise", "revise")):
        costs = [float(r["cost_usd"]) for r in rows if r["step"].startswith(prefix)]
        if costs:
            out[key] = sum(costs) / len(costs)
            out[f"{key}_n"] = len(costs)
    return out if {"intake", "draft", "revise"} <= out.keys() else None


class InputSheet:
    """Writes labelled inputs and remembers each one's absolute cell reference by key."""

    def __init__(self, ws):
        self.ws = ws
        self.row = 4
        self.refs: dict[str, str] = {}

    def section(self, title: str) -> None:
        self.row += 1
        for col in range(1, 5):
            self.ws.cell(self.row, col).fill = SECTION_FILL
        self.ws.cell(self.row, 1, title).font = BOLD
        self.row += 1

    def add(self, key, label, value, fmt, unit="", note="", highlight=False, formula=False):
        ws, r = self.ws, self.row
        ws.cell(r, 1, label).font = BLACK
        c = ws.cell(r, 2, value)
        c.number_format = fmt
        c.font = BLACK if formula else BLUE
        if highlight:
            c.fill = YELLOW
        ws.cell(r, 3, unit).font = BLACK
        ws.cell(r, 4, note).font = Font(name=FONT, italic=True, color="595959")
        ws.cell(r, 4).alignment = Alignment(wrap_text=True, vertical="top")
        for col in range(1, 5):
            ws.cell(r, col).border = THIN
        self.refs[key] = f"Inputs!$B${r}"
        self.row += 1


def build(results_csv: Path, out_path: Path) -> Path:
    bench = benchmark_averages(results_csv)
    wb = Workbook()

    # ------------------------------------------------------------------ Inputs
    ws = wb.active
    ws.title = "Inputs"
    ws["A1"] = "EduNXT agent: cost model inputs"
    ws["A1"].font = TITLE
    ws["A2"] = "Blue = editable input · Yellow = assumption to confirm or value still to fill · Black = formula. All money in USD."
    ws["A2"].font = Font(name=FONT, italic=True)
    ws["A3"] = (
        f"AI costs: averages of {bench['draft_n']} drafts / {bench['revise_n']} revisions / {bench['intake_n']} intake messages "
        f"from out/benchmark/results.csv"
        if bench
        else "AI costs: NOT YET MEASURED. Run `python -m edunxt benchmark`, then `python scripts/build_cost_model.py` (or type the averages from out/benchmark/summary.md into B6:B8)."
    )
    ws["A3"].font = Font(name=FONT, bold=True, color="008000" if bench else "C00000")
    inp = InputSheet(ws)

    inp.section("AI cost per step (Claude API; measured by `python -m edunxt benchmark`)")
    note = "Avg of successful benchmark runs (results.csv)" if bench else "Fill from out/benchmark/summary.md, 'Avg cost (USD)' column"
    inp.add("ai_intake", "AI cost per client message analysed (Intake Agent, Haiku)", bench["intake"] if bench else None, USD4, "USD", note, highlight=not bench)
    inp.add("ai_draft", "AI cost per draft (Research Agent, Opus + web search)", bench["draft"] if bench else None, USD4, "USD", note, highlight=not bench)
    inp.add("ai_revise", "AI cost per founder revision", bench["revise"] if bench else None, USD4, "USD", note, highlight=not bench)

    inp.section("Workflow per enquiry")
    inp.add("msgs_in", "Client messages analysed per enquiry", 2, INT, "messages", "Benchmark cases use 2 (first enquiry + profile). Every later client message is also analysed.")
    inp.add("draft_share", "Share of enquiries that reach a draft", 1.0, PCT, "", "100% = conservative upper bound (every enquirer completes the profile).", highlight=True)
    inp.add("revisions", "Founder revisions per draft", 1, HRS, "revisions", "Benchmark measures 1 revision per case.", highlight=True)
    inp.add("closed_share", "Share of drafts sent while the founder's 24h window is closed", 0.5, PCT, "", "Then the agent sends the `new_case_ready` utility template first.", highlight=True)
    inp.add("templates", "Templates per such draft", 1, INT, "templates", "edunxt/workflow.py `_hold`")
    inp.add("ai_per_case", "AI cost per client (derived)", f"=IF(COUNT(B6:B8)<3,\"not measured\",{inp.refs['msgs_in'].split('!')[1]}*B6+B7+{inp.refs['revisions'].split('!')[1]}*B8)", USD, "USD", "intake x messages + draft + revisions x revision", formula=True)

    inp.section("WhatsApp messages the agent sends (counted from edunxt/workflow.py)")
    inp.add("m_welcome", "Welcome + questionnaire to the client", 1, INT, "per enquiry", "The only message the agent ever sends a client. On Instagram it goes out over Instagram instead.")
    inp.add("m_notify", "Founder notification per client message", 1, INT, "per message", "New-enquiry / forwarded-message notes")
    inp.add("m_draft", "Founder messages per draft", 3, INT, "per draft", "Summary text + University options PDF + Scholarships PDF")
    inp.add("m_revise", "Founder messages per revision", 3, INT, "per revision", "Same three items, new version")
    inp.add("m_other", "Other founder replies per drafted case", 1, INT, "per draft", "e.g. the `approve` confirmation")

    inp.section("Meta pricing, Egypt (USD per delivered message)")
    inp.add("p_service", "Service message (reply inside the 24h window)", 0.0036, USD4, "USD", WA_SOURCE)
    inp.add("p_utility", "Utility template (founder `new_case_ready`)", 0.0036, USD4, "USD", "Volume tiers can lower this by up to 25%; ignored here.")
    inp.add("free_service", "Free service messages per number per month", 1000, INT, "messages", "From 1 Oct 2026; resets monthly, no rollover.")
    inp.add("vat", "VAT added to Meta's invoice", 0.14, PCT, "", "Egyptian VAT 14%. Set 0% if it doesn't apply to your billing entity.", highlight=True)
    inp.add("p_ig", "Instagram message fee", 0, USD4, "USD", IG_SOURCE)
    inp.add("p_marketing", "Marketing template (reference only; the agent sends none)", 0.0644, USD4, "USD", "Not used in the model.")

    inp.section("Fixed monthly costs")
    inp.add("hosting", "Hosting", 0, USD0, "USD / month", "PoC runs on a laptop ($0). A small cloud VM is roughly $5-20/month.", highlight=True)
    inp.add("number", "WhatsApp Cloud API number / access", 0, USD0, "USD / month", "No platform fee from Meta for Cloud API.")

    inp.section("EduNXT business figures")
    inp.add("enrol", "Enquiry to enrolment rate", 0.30, PCT, "", "From EduNXT")
    inp.add("commission", "Commission per student (share of annual university fees)", 0.10, PCT, "", "From EduNXT")
    inp.add("fees", "Average annual university fees", 10000, USD0, "USD / year", "ASSUMPTION for illustration; replace with EduNXT's average.", highlight=True)
    inp.add("list_hours", "Founder time to build one list today", 1.5, HRS, "hours", "From EduNXT: 1-2 hours (midpoint)")
    inp.add("review_hours", "Founder time to review an agent draft", 0.25, HRS, "hours", "ASSUMPTION: 15 minutes to read, tweak and approve.", highlight=True)
    inp.add("current_volume", "Current enquiries per month", 15, INT, "enquiries", "From EduNXT (WhatsApp + Instagram)")

    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 90
    ws.freeze_panes = "A5"
    R = inp.refs

    # ------------------------------------------------------------------ Monthly cost
    mc = wb.create_sheet("Monthly cost")
    mc["A1"] = "Monthly running cost by number of new enquiries"
    mc["A1"].font = TITLE
    mc["A2"] = f'=IF(COUNT(Inputs!B6:B8)<3,"AI cost not yet measured: AI columns show $0 until Inputs!B6:B8 are filled from the benchmark.","AI cost per client: "&TEXT({R["ai_per_case"]},"$0.00"))'
    mc["A2"].font = Font(name=FONT, bold=True, color="C00000")

    cols = [
        ("Enquiries / month", INT),
        ("Drafts", INT),
        ("Client msgs analysed", INT),
        ("WhatsApp service msgs", INT),
        ("Billable service msgs", INT),
        ("Utility templates", INT),
        ("Meta fees (USD)", USD),
        ("AI cost (USD)", USD),
        ("Fixed (USD)", USD),
        ("Total / month (USD)", USD),
        ("Per enquiry (USD)", USD),
        ("Per enrolled student (USD)", USD),
    ]

    def table(top: int, title: str, channel: str) -> tuple[int, int]:
        mc.cell(top, 1, title).font = BOLD
        hr = top + 1
        for j, (name, _) in enumerate(cols, 1):
            c = mc.cell(hr, j, name)
            c.font, c.fill = HEAD, HEAD_FILL
            c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        mc.row_dimensions[hr].height = 32
        first = hr + 1
        for i, vol in enumerate(VOLUMES):
            r = first + i
            welcome = f"A{r}*{R['m_welcome']}" if channel == "whatsapp" else "0"
            ig_fee = "0" if channel == "whatsapp" else f"A{r}*{R['m_welcome']}*{R['p_ig']}"
            values = [
                vol,
                f"=A{r}*{R['draft_share']}",
                f"=A{r}*{R['msgs_in']}",
                f"={welcome}+C{r}*{R['m_notify']}+B{r}*({R['m_draft']}+{R['revisions']}*{R['m_revise']}+{R['m_other']})",
                f"=MAX(0,D{r}-{R['free_service']})",
                f"=B{r}*{R['closed_share']}*{R['templates']}",
                f"=(E{r}*{R['p_service']}+F{r}*{R['p_utility']})*(1+{R['vat']})+{ig_fee}",
                f"=C{r}*N({R['ai_intake']})+B{r}*(N({R['ai_draft']})+{R['revisions']}*N({R['ai_revise']}))",
                f"={R['hosting']}+{R['number']}",
                f"=G{r}+H{r}+I{r}",
                f"=IF(A{r}>0,J{r}/A{r},0)",
                f"=IF(A{r}*{R['enrol']}>0,J{r}/(A{r}*{R['enrol']}),0)",
            ]
            for j, v in enumerate(values, 1):
                c = mc.cell(r, j, v)
                c.number_format = cols[j - 1][1]
                c.font = BLUE if j == 1 else BLACK
                c.border = THIN
        return first, first + len(VOLUMES) - 1

    wa_first, wa_last = table(4, "WhatsApp: client enquires on WhatsApp; agent replies on WhatsApp; founders review on WhatsApp", "whatsapp")
    ig_top = wa_last + 3
    ig_first, ig_last = table(
        ig_top,
        "Instagram: client enquires by Instagram DM; welcome sent over the Instagram Messaging API (free); founders still review on WhatsApp",
        "instagram",
    )
    note_row = ig_last + 2
    notes = [
        "WhatsApp service messages = welcome (WhatsApp only) + one founder note per client message + founder messages per draft and revision.",
        "Only service messages above the free 1,000 per number per month are billed, so at today's volume Meta fees are close to $0.",
        "Instagram: the agent's single client message is inside Instagram's 24h window, so the window rule doesn't limit it. "
        "Integrating Instagram needs Meta App Review (instagram_business_manage_messages); until then founders start Instagram cases with `new ...` on WhatsApp.",
        "Volumes in column A are editable.",
    ]
    for k, text in enumerate(notes):
        mc.cell(note_row + k, 1, text).font = Font(name=FONT, italic=True, color="595959")

    for j, w in enumerate([12, 9, 12, 13, 12, 11, 13, 13, 11, 15, 13, 15], 1):
        mc.column_dimensions[get_column_letter(j)].width = w

    chart = LineChart()
    chart.title = "Total monthly cost (USD)"
    chart.y_axis.title = "USD / month"
    chart.x_axis.title = "New enquiries / month"
    chart.height, chart.width = 8, 18
    chart.add_data(Reference(mc, min_col=10, min_row=wa_first, max_row=wa_last), titles_from_data=False)
    chart.add_data(Reference(mc, min_col=10, min_row=ig_first, max_row=ig_last), titles_from_data=False)
    chart.series[0].tx = SeriesLabel(v="WhatsApp")
    chart.series[1].tx = SeriesLabel(v="Instagram")
    chart.set_categories(Reference(mc, min_col=1, min_row=wa_first, max_row=wa_last))
    mc.add_chart(chart, f"N{wa_first - 1}")
    mc.freeze_panes = "B4"

    # ------------------------------------------------------------------ Value
    va = wb.create_sheet("Value")
    va["A1"] = "What the cost buys: time saved and commission at stake (WhatsApp table)"
    va["A1"].font = TITLE
    va["A2"] = "Commission uses the ASSUMED average annual fees on Inputs; replace it before quoting revenue."
    va["A2"].font = Font(name=FONT, italic=True, color="C00000")
    vcols = [
        ("Enquiries / month", INT),
        ("Enrolled students", HRS),
        ("Commission / month (USD)", USD0),
        ("Agent cost / month (USD)", USD),
        ("Agent cost as % of commission", PCT),
        ("Founder hours on lists today", HRS),
        ("Founder hours with agent", HRS),
        ("Founder hours saved / month", HRS),
    ]
    for j, (name, _) in enumerate(vcols, 1):
        c = va.cell(4, j, name)
        c.font, c.fill = HEAD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    va.row_dimensions[4].height = 32
    for i in range(len(VOLUMES) + 1):
        r = 5 + i
        src = wa_first + i - 1  # row 5 is today's volume; the rest mirror the WhatsApp table
        vol = f"={R['current_volume']}" if i == 0 else f"='Monthly cost'!A{src}"
        cost = (
            f"=N({R['ai_per_case']})*A{r}*{R['draft_share']}"  # today's 15: well inside the free allowance
            if i == 0
            else f"='Monthly cost'!J{src}"
        )
        values = [
            vol,
            f"=A{r}*{R['enrol']}",
            f"=B{r}*{R['fees']}*{R['commission']}",
            cost,
            f"=IF(C{r}>0,D{r}/C{r},0)",
            f"=A{r}*{R['draft_share']}*{R['list_hours']}",
            f"=A{r}*{R['draft_share']}*{R['review_hours']}",
            f"=F{r}-G{r}",
        ]
        for j, v in enumerate(values, 1):
            c = va.cell(r, j, v)
            c.number_format = vcols[j - 1][1]
            c.font = GREEN if (j == 1 or (j == 4 and i > 0)) else BLACK
            c.border = THIN
    va.cell(5, 9, "← EduNXT today").font = BOLD
    for j in range(1, 9):
        va.column_dimensions[get_column_letter(j)].width = 15
    va.column_dimensions["I"].width = 16

    # ------------------------------------------------------------------ Sensitivity
    se = wb.create_sheet("Sensitivity")
    se["A1"] = "Total monthly cost (WhatsApp) if the AI cost per client were..."
    se["A1"].font = TITLE
    se["A2"] = "Use this before the benchmark has run, or to see how much the AI cost matters. Meta fees and fixed costs come from 'Monthly cost'."
    se["A2"].font = Font(name=FONT, italic=True)
    se.cell(4, 1, "Enquiries / month").font = HEAD
    se.cell(4, 1).fill = HEAD_FILL
    for k, s in enumerate(SCENARIOS):
        c = se.cell(4, 2 + k, s)
        c.number_format = '"$"0.00" / client"'
        c.font, c.fill = Font(name=FONT, bold=True, color="FFFFFF"), HEAD_FILL
    for i in range(len(VOLUMES)):
        r, src = 5 + i, wa_first + i
        se.cell(r, 1, f"='Monthly cost'!A{src}").font = GREEN
        se.cell(r, 1).number_format = INT
        for k in range(len(SCENARIOS)):
            col = get_column_letter(2 + k)
            c = se.cell(r, 2 + k, f"='Monthly cost'!G{src}+'Monthly cost'!I{src}+A{r}*{R['draft_share']}*{col}$4")
            c.number_format = USD0
            c.border = THIN
    se.column_dimensions["A"].width = 18
    for k in range(len(SCENARIOS)):
        se.column_dimensions[get_column_letter(2 + k)].width = 16
    se["A4"].comment = Comment("Header values (row 4) are editable scenarios.", "EduNXT")

    # ------------------------------------------------------------------ Sources
    so = wb.create_sheet("Sources")
    lines = [
        ("Sources and notes", TITLE),
        ("WhatsApp pricing (Egypt): " + WA_SOURCE, BLACK),
        ("Instagram Messaging API: " + IG_SOURCE, BLACK),
        ("Claude API prices used by the cost meter (edunxt/agents/llm.py): Opus 5.5 $4 / $20 per million input / output tokens; Haiku 4.5 $1 / $5; web search $10 per 1,000.", BLACK),
        ("Message counts per case: edunxt/workflow.py (_handle_client, _deliver, _hold).", BLACK),
        ("Business figures (15 enquiries/month, 30% enrolment, 10% commission, 1-2 h per list, 12-24 h reply time): EduNXT founders.", BLACK),
        ("Generated by scripts/build_cost_model.py. Re-run it after a benchmark to refresh the AI-cost inputs.", BLACK),
    ]
    for k, (text, font) in enumerate(lines, 1):
        c = so.cell(k, 1, text)
        c.font = font
        c.alignment = Alignment(wrap_text=True, vertical="top")
    so.column_dimensions["A"].width = 140

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--results", type=Path, default=ROOT / "out" / "benchmark" / "results.csv")
    ap.add_argument("--out", type=Path, default=ROOT / "docs" / "cost_model.xlsx")
    args = ap.parse_args()
    path = build(args.results, args.out)
    print(f"Wrote {path}" + ("" if benchmark_averages(args.results) else " (AI-cost inputs left blank: no benchmark results found)"))


if __name__ == "__main__":
    main()
